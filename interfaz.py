"""Interfaz para consultar las marcas de asistencia del equipo Dahua.

Ejecutar con:  py interfaz.py
"""

import calendar
import threading
import tkinter as tk
from collections import defaultdict
from datetime import date, datetime, timedelta
from time import mktime
from tkinter import filedialog, messagebox, ttk

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from requests.auth import HTTPDigestAuth

# --- Configuración ----------------------------------------------------------

IP = "192.168.3.1"
USUARIO = "admin"
CLAVE = "qwerty1234"
MAX_REGISTROS = 100  # registros que se piden al equipo en cada consulta

# Códigos de AttendanceState según el Excel exportado desde el equipo.
ESTADOS_ASISTENCIA = {
    "0": "-",
    "1": "Ingreso",
    "2": "Escape",
    "3": "Intrusión",
    "4": "Salida",
}
ESTADOS_VALIDOS = {"Ingreso", "Salida", "Intrusión", "Escape"}
METODOS = {"15": "Cara"}
RESULTADOS = {"1": "Exitoso", "0": "Fallido"}

MESES = [
    "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio",
    "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre",
]
DIAS = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]
DIAS_CORTOS = [d[:3].capitalize() for d in DIAS]  # Lun, Mar, Mié, ...
AZUL = "#1F4E78"
AYUDA_RESUMEN = (
    "Selecciona un usuario para ver sus observaciones y sus marcas. "
    "Rojo = faltan marcas (esas horas no se cuentan). "
    "Amarillo = marcas mal hechas que se corrigieron, conviene revisarlas."
)
AYUDA_SEMANA = AYUDA_RESUMEN + " ⚠ = ese día tiene observaciones."

# Emojis de Slack para el resumen copiado (mismos colores que la tabla).
ICONOS_SLACK = {
    "Completo": ":large_green_circle:",
    "Revisar": ":large_yellow_circle:",
    "Incompleto": ":red_circle:",
}

COLUMNAS_REGISTROS = [
    ("fecha", "Fecha", 90, "center"),
    ("hora", "Hora", 90, "center"),
    ("id", "ID", 60, "center"),
    ("nombre", "Nombre", 220, "w"),
    ("marca", "Estado de Asistencia", 150, "center"),
    ("metodo", "Método", 100, "center"),
    ("resultado", "Estado", 100, "center"),
]

COLUMNAS_RESUMEN = [
    ("id", "ID", 50, "center"),
    ("nombre", "Nombre", 170, "w"),
    ("ingreso", "Ingreso", 75, "center"),
    ("salida", "Salida", 75, "center"),
    ("pausas", "Escapes", 70, "center"),
    ("fuera", "Tiempo fuera", 95, "center"),
    ("trabajado", "Horas laboradas", 115, "center"),
    ("estado", "Estado", 90, "center"),
    ("obs", "Observaciones", 380, "w"),
]


# --- Consulta al equipo -----------------------------------------------------

def inicio_del_dia(dia):
    return int(mktime(datetime(dia.year, dia.month, dia.day).timetuple()))


def consultar_registros(dia):
    """Devuelve todas las marcas del día, ordenadas por hora."""
    respuesta = requests.get(
        f"http://{IP}/cgi-bin/recordFinder.cgi",
        params={
            "action": "find",
            "name": "AccessControlCardRec",
            "StartTime": inicio_del_dia(dia),
            "EndTime": inicio_del_dia(dia + timedelta(days=1)),
            "count": MAX_REGISTROS,
        },
        auth=HTTPDigestAuth(USUARIO, CLAVE),
        timeout=30,
    )
    respuesta.raise_for_status()

    # Convierte las líneas records[0].Campo=valor en diccionarios.
    crudos = {}
    for linea in respuesta.text.splitlines():
        clave, igual, valor = linea.partition("=")
        if not igual or not clave.startswith("records["):
            continue
        indice, _, campo = clave[len("records["):].partition("].")
        if indice.isdigit() and campo:
            crudos.setdefault(int(indice), {})[campo] = valor.strip()

    registros = []
    for crudo in crudos.values():
        try:
            fecha = datetime.fromtimestamp(int(crudo.get("CreateTime", "")))
        except (ValueError, OSError):
            continue
        metodo = crudo.get("Method", "")
        resultado = crudo.get("Status", "")
        registros.append({
            "fecha": fecha,
            "id": crudo.get("UserID", ""),
            "nombre": crudo.get("CardName", ""),
            "marca": ESTADOS_ASISTENCIA.get(crudo.get("AttendanceState", ""), "-"),
            "metodo": METODOS.get(metodo, metodo),
            "resultado": RESULTADOS.get(resultado, resultado),
        })

    registros.sort(key=lambda r: r["fecha"])
    return registros


def mensaje_error(error):
    if isinstance(error, (requests.ConnectionError, requests.Timeout)):
        return (
            f"No se pudo conectar con el equipo ({IP}).\n"
            "Revisa que esté encendido y en la misma red."
        )
    respuesta = getattr(error, "response", None)
    if respuesta is not None and respuesta.status_code == 401:
        return "El equipo rechazó el usuario o la clave."
    return f"Ocurrió un error al consultar el equipo:\n{error}"


# --- Cálculo de horas -------------------------------------------------------

def hora(fecha):
    return fecha.strftime("%H:%M")


def duracion(tiempo):
    minutos = int(tiempo.total_seconds() // 60)
    return f"{minutos // 60}:{minutos % 60:02d}"


def calcular_jornada(marcas):
    """Calcula la jornada de un usuario a partir de sus marcas en orden.

    La jornada esperada es: un Ingreso, luego 0..n pares Escape -> Intrusión,
    y al final una Salida. Solo se suman los tramos que tienen marca de inicio
    y de fin; si falta una marca, ese tramo no se cuenta y se indica.
    """
    trabajado = fuera = timedelta()
    estado = "sin_ingreso"  # sin_ingreso | dentro | pausa | terminado
    desde = None            # hora en que empezó el tramo actual
    ingreso = salida = None
    pausas = 0
    faltas, revisar, sin_estado = [], [], []

    for marca in marcas:
        t, tipo, h = marca["fecha"], marca["marca"], hora(marca["fecha"])

        if tipo not in ESTADOS_VALIDOS:
            sin_estado.append(h)
            continue

        if estado == "terminado":
            if tipo in ("Escape", "Salida"):
                revisar.append(f"{tipo} {h} después de la Salida, no se tomó en cuenta")
                continue
            # Volvió a entrar: esa Salida en realidad fue un Escape.
            revisar.append(f"Salida {hora(salida)} seguida de más marcas, se tomó como Escape")
            estado, desde, salida = "pausa", salida, None
            pausas += 1

        if tipo == "Ingreso":
            if estado == "sin_ingreso":
                ingreso, desde, estado = t, t, "dentro"
                continue
            if estado == "dentro":
                revisar.append(f"Ingreso {h} de más, no se tomó en cuenta")
                continue
            revisar.append(f"Ingreso {h} después de un Escape, se tomó como Intrusión")
            tipo = "Intrusión"

        if tipo == "Intrusión":
            if estado == "pausa":
                fuera += t - desde
            elif estado == "dentro":
                faltas.append(
                    f"Falta Escape antes de la Intrusión {h} "
                    f"(tramo {hora(desde)}–{h} sin contar)"
                )
            else:
                faltas.append(f"Falta Ingreso (primera marca: Intrusión {h})")
            desde, estado = t, "dentro"

        elif tipo == "Escape":
            if estado == "dentro":
                trabajado += t - desde
            elif estado == "pausa":
                faltas.append(f"Falta Intrusión entre los Escapes de {hora(desde)} y {h}")
            else:
                faltas.append(f"Falta Ingreso (primera marca: Escape {h})")
            desde, estado = t, "pausa"
            pausas += 1

        else:  # Salida
            if estado == "dentro":
                trabajado += t - desde
            elif estado == "pausa":
                faltas.append(f"Falta Intrusión después del Escape {hora(desde)}")
            else:
                faltas.append(f"Falta Ingreso (primera marca: Salida {h})")
            salida, estado = t, "terminado"

    if estado == "dentro":
        faltas.append(f"Falta Salida (tramo desde {hora(desde)} sin contar)")
    elif estado == "pausa":
        faltas.append(f"Falta Intrusión y Salida después del Escape {hora(desde)}")
    elif estado == "sin_ingreso":
        faltas.append("Sin marcas de asistencia")

    observaciones = faltas + revisar
    if sin_estado:
        observaciones.append(f"Marcas sin estado (no se cuentan): {', '.join(sin_estado)}")

    return {
        "ingreso": ingreso,
        "salida": salida,
        "pausas": pausas,
        "fuera": fuera,
        "trabajado": trabajado,
        "estado": "Incompleto" if faltas else "Revisar" if revisar else "Completo",
        "observaciones": observaciones,
    }


def resumir(registros):
    """Una fila por usuario con su jornada calculada (solo marcas exitosas)."""
    por_usuario = defaultdict(list)
    for r in registros:
        if r["resultado"] == "Exitoso":
            por_usuario[r["id"]].append(r)

    filas = []
    for uid, marcas in por_usuario.items():
        fila = calcular_jornada(marcas)  # los registros ya vienen ordenados
        fila["id"] = uid
        fila["nombre"] = next((m["nombre"] for m in marcas if m["nombre"]), "")
        filas.append(fila)

    filas.sort(key=lambda f: f["nombre"].lower())
    return filas


def resumir_semana(dias):
    """Una fila por usuario con su jornada de cada día y los totales de la semana.

    `dias` es una lista de (fecha, registros de ese día).
    """
    usuarios = {}
    for dia, registros in dias:
        for jornada in resumir(registros):
            u = usuarios.setdefault(jornada["id"], {"id": jornada["id"], "nombre": "", "dias": {}})
            u["nombre"] = u["nombre"] or jornada["nombre"]
            u["dias"][dia] = jornada

    filas = list(usuarios.values())
    for u in filas:
        jornadas = list(u["dias"].values())
        estados = {j["estado"] for j in jornadas}
        u["trabajado"] = sum((j["trabajado"] for j in jornadas), timedelta())
        u["fuera"] = sum((j["fuera"] for j in jornadas), timedelta())
        u["estado"] = next(e for e in ("Incompleto", "Revisar", "Completo") if e in estados)
        u["observaciones"] = [
            f"{dia_corto(dia)}: {obs}"
            for dia, j in u["dias"].items()
            for obs in j["observaciones"]
        ]

    filas.sort(key=lambda f: f["nombre"].lower())
    return filas


def lunes(dia):
    return dia - timedelta(days=dia.weekday())


def dia_corto(dia):
    return f"{DIAS_CORTOS[dia.weekday()]} {dia.day}"


def texto_periodo(inicio, semanal):
    if not semanal:
        return inicio.strftime("%d-%m-%Y")
    return f"{inicio:%d-%m-%Y} al {inicio + timedelta(days=6):%d-%m-%Y}"


def columnas_semana(inicio):
    """Columnas del resumen semanal; cada día lleva su fecha en el título."""
    dias = [(f"d{i}", dia_corto(inicio + timedelta(days=i)), 58, "center") for i in range(7)]
    return [
        ("id", "ID", 45, "center"),
        ("nombre", "Nombre", 160, "w"),
        *dias,
        ("dias", "Días", 45, "center"),
        ("fuera", "Tiempo fuera", 90, "center"),
        ("trabajado", "Total horas", 85, "center"),
        ("estado", "Estado", 85, "center"),
        ("obs", "Observaciones", 240, "w"),
    ]


def fila_registro(r):
    return (
        r["fecha"].strftime("%d-%m-%Y"),
        r["fecha"].strftime("%H:%M:%S"),
        r["id"],
        r["nombre"],
        r["marca"],
        r["metodo"],
        r["resultado"],
    )


def fila_resumen(f):
    return (
        f["id"],
        f["nombre"],
        hora(f["ingreso"]) if f["ingreso"] else "—",
        hora(f["salida"]) if f["salida"] else "—",
        f["pausas"],
        duracion(f["fuera"]),
        duracion(f["trabajado"]),
        f["estado"],
        "; ".join(f["observaciones"]),
    )


def fila_semana(f, inicio):
    horas_por_dia = []
    for i in range(7):
        jornada = f["dias"].get(inicio + timedelta(days=i))
        if jornada is None:
            horas_por_dia.append("—")
        else:
            aviso = "" if jornada["estado"] == "Completo" else " ⚠"
            horas_por_dia.append(duracion(jornada["trabajado"]) + aviso)
    return (
        f["id"],
        f["nombre"],
        *horas_por_dia,
        len(f["dias"]),
        duracion(f["fuera"]),
        duracion(f["trabajado"]),
        f["estado"],
        "; ".join(f["observaciones"]),
    )


# --- Texto para Slack -------------------------------------------------------

def plural(n, palabra):
    return f"{n} {palabra}{'' if n == 1 else 's'}"


def texto_slack(resumen, inicio, semanal):
    """Resumen con formato de Slack (*negrita* y :emojis:), listo para pegar."""
    if semanal:
        titulo = f"Asistencia de la semana del {texto_periodo(inicio, True)}"
    else:
        titulo = f"Asistencia del {DIAS[inicio.weekday()]} {inicio:%d-%m-%Y}"

    cuenta = [plural(len(resumen), "persona")]
    incompletos = sum(f["estado"] == "Incompleto" for f in resumen)
    revisar = sum(f["estado"] == "Revisar" for f in resumen)
    if incompletos:
        cuenta.append(f"{incompletos} con marcas faltantes")
    if revisar:
        cuenta.append(f"{revisar} para revisar")

    lineas = [f"*{titulo}*", " · ".join(cuenta), ""]
    for f in resumen:
        detalle = f"{duracion(f['trabajado'])} h"
        if semanal:
            detalle += f" en {plural(len(f['dias']), 'día')}"
        elif f["ingreso"] or f["salida"]:
            entrada = hora(f["ingreso"]) if f["ingreso"] else "—"
            salida = hora(f["salida"]) if f["salida"] else "—"
            detalle += f" ({entrada} a {salida})"
        nombre = f["nombre"] or f"ID {f['id']}"
        lineas.append(f"{ICONOS_SLACK[f['estado']]} *{nombre}*: {detalle}")
        lineas.extend(f"      • {obs}" for obs in f["observaciones"])

    lineas += [
        "",
        f"{ICONOS_SLACK['Completo']} completo   {ICONOS_SLACK['Revisar']} revisar   "
        f"{ICONOS_SLACK['Incompleto']} faltan marcas (esas horas no se cuentan)",
    ]
    return "\n".join(lineas)


# --- Excel ------------------------------------------------------------------

def guardar_excel(archivo, registros, columnas_resumen, filas_resumen):
    wb = Workbook()
    hojas = [
        (wb.active, "Registros", COLUMNAS_REGISTROS, [fila_registro(r) for r in registros]),
        (wb.create_sheet(), "Resumen", columnas_resumen, filas_resumen),
    ]
    for hoja, titulo, columnas, filas in hojas:
        hoja.title = titulo
        hoja.append([c[1] for c in columnas])
        for fila in filas:
            hoja.append(fila)

        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = hoja.dimensions
        for celda in hoja[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill(fill_type="solid", fgColor=AZUL[1:])
            celda.alignment = Alignment(horizontal="center")
        for columna in hoja.columns:
            ancho = max(len(str(c.value or "")) for c in columna)
            hoja.column_dimensions[columna[0].column_letter].width = min(ancho + 3, 60)

    wb.save(archivo)


# --- Selector de fecha ------------------------------------------------------

class Calendario(tk.Toplevel):
    """Ventana pequeña con el mes para elegir un día (o la semana de ese día)."""

    def __init__(self, selector, fecha, al_elegir, semanal=False):
        super().__init__(selector)
        self.title("Elegir semana" if semanal else "Elegir fecha")
        self.resizable(False, False)
        self.transient(selector.winfo_toplevel())
        self.al_elegir = al_elegir
        # Días resaltados: el elegido, o los 7 de la semana elegida.
        self.marcados = {fecha + timedelta(days=i) for i in range(7 if semanal else 1)}
        self.anio, self.mes = fecha.year, fecha.month

        cabecera = ttk.Frame(self, padding=(8, 8, 8, 4))
        cabecera.pack(fill="x")
        ttk.Button(cabecera, text="◀", width=3, command=lambda: self.cambiar_mes(-1)).pack(side="left")
        ttk.Button(cabecera, text="▶", width=3, command=lambda: self.cambiar_mes(1)).pack(side="right")
        self.titulo = ttk.Label(cabecera, anchor="center", font=("Segoe UI", 10, "bold"))
        self.titulo.pack(fill="x", expand=True)

        self.cuadricula = ttk.Frame(self, padding=(8, 0, 8, 8))
        self.cuadricula.pack()
        self.dibujar()

        x = selector.winfo_rootx()
        y = selector.winfo_rooty() + selector.winfo_height() + 2
        self.geometry(f"+{x}+{y}")
        self.bind("<Escape>", lambda _e: self.destroy())
        self.wait_visibility()
        self.grab_set()
        self.focus_set()

    def cambiar_mes(self, delta):
        indice = self.mes - 1 + delta
        self.anio += indice // 12
        self.mes = indice % 12 + 1
        self.dibujar()

    def dibujar(self):
        for widget in self.cuadricula.winfo_children():
            widget.destroy()
        self.titulo.config(text=f"{MESES[self.mes - 1]} {self.anio}")

        for col, nombre in enumerate(["Lu", "Ma", "Mi", "Ju", "Vi", "Sá", "Do"]):
            ttk.Label(self.cuadricula, text=nombre, width=4, anchor="center").grid(row=0, column=col)

        hoy = date.today()
        for fila, semana in enumerate(calendar.monthcalendar(self.anio, self.mes), start=1):
            for col, numero in enumerate(semana):
                if not numero:
                    continue
                dia = date(self.anio, self.mes, numero)
                boton = tk.Button(
                    self.cuadricula, text=numero, width=3, relief="flat",
                    command=lambda d=dia: self.elegir(d),
                )
                if dia in self.marcados:
                    boton.config(bg=AZUL, fg="white", activebackground=AZUL, activeforeground="white")
                elif dia == hoy:
                    boton.config(fg=AZUL, font=("Segoe UI", 9, "bold"))
                if dia > hoy:
                    boton.config(state="disabled")
                boton.grid(row=fila, column=col, padx=1, pady=1)

    def elegir(self, dia):
        self.al_elegir(dia)
        self.destroy()


class SelectorFecha(ttk.Frame):
    """Campo de fecha con anterior/siguiente y calendario.

    En modo semanal el periodo es la semana (lunes a domingo) del día elegido
    y las flechas avanzan de semana en semana.
    """

    def __init__(self, padre):
        super().__init__(padre)
        self.texto = tk.StringVar()
        self.semanal = False

        ttk.Button(self, text="◀", width=3, command=lambda: self.mover(-1)).pack(side="left")
        self.campo = ttk.Entry(self, textvariable=self.texto, width=12, state="readonly", justify="center")
        self.campo.pack(side="left", padx=2)
        self.campo.bind("<Button-1>", lambda _e: self.abrir_calendario())
        self.btn_siguiente = ttk.Button(self, text="▶", width=3, command=lambda: self.mover(1))
        self.btn_siguiente.pack(side="left")
        ttk.Button(self, text="▼", width=3, command=self.abrir_calendario).pack(side="left", padx=(4, 0))

        self.poner(date.today())

    def poner(self, dia):
        hoy = date.today()
        self.dia = min(dia, hoy)  # se recuerda para volver al mismo día desde la semana
        self.fecha = lunes(self.dia) if self.semanal else self.dia  # inicio del periodo
        self.texto.set(texto_periodo(self.fecha, self.semanal))
        ultimo = lunes(hoy) if self.semanal else hoy
        self.btn_siguiente.state(["disabled"] if self.fecha >= ultimo else ["!disabled"])

    def mover(self, sentido):
        self.poner(self.dia + timedelta(days=sentido * (7 if self.semanal else 1)))

    def cambiar_modo(self, semanal):
        self.semanal = semanal
        self.campo.config(width=23 if semanal else 12)
        self.poner(self.dia)

    def abrir_calendario(self):
        Calendario(self, self.fecha, self.poner, self.semanal)


# --- Ventana principal ------------------------------------------------------

def poner_columnas(tabla, columnas):
    tabla.configure(columns=[c[0] for c in columnas])
    for clave, titulo, ancho, ancla in columnas:
        tabla.heading(clave, text=titulo, anchor=ancla)
        tabla.column(clave, width=ancho, minwidth=40, anchor=ancla,
                     stretch=clave in ("nombre", "obs"))


def crear_tabla(padre, columnas):
    marco = ttk.Frame(padre)
    marco.pack(fill="both", expand=True)
    tabla = ttk.Treeview(marco, show="headings", selectmode="browse")
    poner_columnas(tabla, columnas)
    barra = ttk.Scrollbar(marco, orient="vertical", command=tabla.yview)
    tabla.configure(yscrollcommand=barra.set)
    barra.pack(side="right", fill="y")
    tabla.pack(side="left", fill="both", expand=True)
    return tabla


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Asistencia - Control de acceso")
        self.geometry("1150x720")
        self.minsize(960, 550)

        # Datos cargados (cambian solo al presionar Generar).
        self.inicio = None     # primer día del periodo
        self.semanal = False
        self.registros = []    # marcas con estado de asistencia válido
        self.resumen = []      # una fila por usuario (del día o de la semana)
        self.columnas_resumen = COLUMNAS_RESUMEN
        self.filas_resumen = []
        self.usuarios = {}     # texto del filtro -> ID de usuario

        estilo = ttk.Style(self)
        estilo.configure("Treeview", rowheight=24)
        estilo.configure("Treeview.Heading", font=("Segoe UI", 9, "bold"))

        # Barra superior
        barra = ttk.Frame(self, padding=10)
        barra.pack(fill="x")
        self.modo = tk.StringVar(value="dia")
        for texto, valor in (("Día", "dia"), ("Semana", "semana")):
            ttk.Radiobutton(
                barra, text=texto, value=valor, variable=self.modo,
                command=lambda: self.selector.cambiar_modo(self.modo.get() == "semana"),
            ).pack(side="left", padx=(0, 6))
        self.selector = SelectorFecha(barra)
        self.selector.pack(side="left", padx=(6, 12))
        self.btn_generar = ttk.Button(barra, text="Generar", command=self.generar)
        self.btn_generar.pack(side="left")
        self.btn_exportar = ttk.Button(barra, text="Exportar a Excel", command=self.exportar)
        self.btn_exportar.pack(side="left", padx=(6, 0))
        self.btn_copiar = ttk.Button(barra, text="Copiar resumen", command=self.copiar_resumen)
        self.btn_copiar.pack(side="left", padx=6)
        for boton in (self.btn_exportar, self.btn_copiar):
            boton.state(["disabled"])

        ttk.Label(barra, text="Usuario:").pack(side="left", padx=(18, 6))
        self.filtro = ttk.Combobox(barra, state="readonly", width=24, values=["Todos"])
        self.filtro.set("Todos")
        self.filtro.pack(side="left")
        self.filtro.bind("<<ComboboxSelected>>", lambda _e: self.llenar_registros())

        # Barra de estado
        self.estado = tk.StringVar(value="Elige una fecha y presiona Generar.")
        ttk.Label(self, textvariable=self.estado, padding=(10, 0, 10, 8)).pack(side="bottom", fill="x")

        # Tablas
        paneles = ttk.PanedWindow(self, orient="vertical")
        paneles.pack(fill="both", expand=True, padx=10, pady=(0, 6))

        self.marco_registros = ttk.LabelFrame(paneles, text="Registros", padding=6)
        self.tabla_registros = crear_tabla(self.marco_registros, COLUMNAS_REGISTROS)
        paneles.add(self.marco_registros, weight=3)

        self.marco_resumen = ttk.LabelFrame(paneles, text="Resumen por usuario", padding=6)
        self.detalle = tk.StringVar(value=AYUDA_RESUMEN)
        detalle = ttk.Label(self.marco_resumen, textvariable=self.detalle, justify="left")
        detalle.pack(side="bottom", fill="x", pady=(6, 0))
        detalle.bind("<Configure>", lambda e: detalle.config(wraplength=e.width - 10))
        self.tabla_resumen = crear_tabla(self.marco_resumen, COLUMNAS_RESUMEN)
        self.tabla_resumen.tag_configure("Incompleto", background="#F8D7DA")
        self.tabla_resumen.tag_configure("Revisar", background="#FFF3CD")
        self.tabla_resumen.bind("<<TreeviewSelect>>", self.al_elegir_resumen)
        paneles.add(self.marco_resumen, weight=2)

    # Consulta en segundo plano para que la ventana no se congele.
    def generar(self):
        inicio, semanal = self.selector.fecha, self.selector.semanal
        # Un día por consulta (el equipo limita la cantidad de registros);
        # los días de la semana que aún no llegan no se consultan.
        dias = [inicio + timedelta(days=i) for i in range(7 if semanal else 1)]
        dias = [d for d in dias if d <= date.today()]
        self.btn_generar.state(["disabled"])
        self.config(cursor="watch")
        self.estado.set(f"Consultando el equipo ({IP})…")
        resultado = {}

        def consultar():
            try:
                resultado["dias"] = [(d, consultar_registros(d)) for d in dias]
            except Exception as error:  # se le muestra al usuario
                resultado["error"] = error

        hilo = threading.Thread(target=consultar, daemon=True)
        hilo.start()
        self.esperar(hilo, resultado, inicio, semanal)

    def esperar(self, hilo, resultado, inicio, semanal):
        if hilo.is_alive():
            self.after(100, self.esperar, hilo, resultado, inicio, semanal)
            return
        self.btn_generar.state(["!disabled"])
        self.config(cursor="")
        if "error" in resultado:
            self.estado.set("No se pudo obtener la información.")
            messagebox.showerror("Error", mensaje_error(resultado["error"]), parent=self)
            return
        self.mostrar(inicio, semanal, resultado["dias"])

    def mostrar(self, inicio, semanal, dias):
        self.inicio, self.semanal = inicio, semanal
        todos = [r for _, registros in dias for r in registros]
        self.registros = [r for r in todos if r["marca"] in ESTADOS_VALIDOS]
        if semanal:
            self.resumen = resumir_semana(dias)
            self.columnas_resumen = columnas_semana(inicio)
            self.filas_resumen = [fila_semana(f, inicio) for f in self.resumen]
        else:
            self.resumen = resumir(todos)
            self.columnas_resumen = COLUMNAS_RESUMEN
            self.filas_resumen = [fila_resumen(f) for f in self.resumen]
        self.usuarios = {f["nombre"] or f"ID {f['id']}": f["id"] for f in self.resumen}

        self.filtro["values"] = ["Todos", *self.usuarios]
        if self.filtro.get() not in self.usuarios:
            self.filtro.set("Todos")

        periodo = texto_periodo(inicio, semanal)
        self.marco_registros.config(text=f"Registros del {periodo}")
        self.marco_resumen.config(text=f"Resumen por usuario del {periodo}")
        self.llenar_registros()
        self.llenar_resumen()
        for boton in (self.btn_exportar, self.btn_copiar):
            boton.state(["!disabled"] if self.resumen else ["disabled"])

        if not todos:
            self.estado.set(f"No hay marcas para el {periodo}.")
            return
        texto = f"{periodo}: {len(self.registros)} marcas de {len(self.resumen)} usuarios."
        incompletos = sum(f["estado"] == "Incompleto" for f in self.resumen)
        if incompletos:
            texto += f" {incompletos} con marcas faltantes."
        llenos = [d for d, registros in dias if len(registros) >= MAX_REGISTROS]
        if llenos:
            cuando = f" el {', '.join(dia_corto(d) for d in llenos)}" if semanal else ""
            texto += (
                f" Atención: el equipo devolvió el máximo de {MAX_REGISTROS} "
                f"registros{cuando}, puede que falten marcas."
            )
        self.estado.set(texto)

    def llenar_registros(self):
        tabla = self.tabla_registros
        tabla.delete(*tabla.get_children())
        uid = self.usuarios.get(self.filtro.get())  # None = Todos
        for r in self.registros:
            if uid is None or r["id"] == uid:
                tabla.insert("", "end", values=fila_registro(r))

    def llenar_resumen(self):
        tabla = self.tabla_resumen
        tabla.delete(*tabla.get_children())
        poner_columnas(tabla, self.columnas_resumen)
        for i, (f, valores) in enumerate(zip(self.resumen, self.filas_resumen)):
            tabla.insert("", "end", iid=str(i), values=valores, tags=(f["estado"],))
        self.detalle.set(AYUDA_SEMANA if self.semanal else AYUDA_RESUMEN)

    def al_elegir_resumen(self, _evento):
        seleccion = self.tabla_resumen.selection()
        if not seleccion:
            return
        fila = self.resumen[int(seleccion[0])]
        sin_obs = f"{'Semana' if self.semanal else 'Jornada'} completa, sin observaciones."
        self.detalle.set(f"{fila['nombre']}: {'; '.join(fila['observaciones']) or sin_obs}")
        etiqueta = next((k for k, v in self.usuarios.items() if v == fila["id"]), "Todos")
        self.filtro.set(etiqueta)
        self.llenar_registros()

    def exportar(self):
        archivo = filedialog.asksaveasfilename(
            parent=self,
            title="Guardar Excel",
            defaultextension=".xlsx",
            filetypes=[("Libro de Excel", "*.xlsx")],
            initialfile=f"asistencia_{texto_periodo(self.inicio, self.semanal).replace(' ', '_')}.xlsx",
        )
        if not archivo:
            return
        try:
            guardar_excel(archivo, self.registros, self.columnas_resumen, self.filas_resumen)
        except PermissionError:
            messagebox.showerror(
                "No se pudo guardar",
                "El archivo está abierto en otro programa. Ciérralo e intenta de nuevo.",
                parent=self,
            )
            return
        self.estado.set(f"Excel guardado: {archivo}")

    def copiar_resumen(self):
        self.clipboard_clear()
        self.clipboard_append(texto_slack(self.resumen, self.inicio, self.semanal))
        self.estado.set("Resumen copiado. Pégalo en Slack con Ctrl+V.")


if __name__ == "__main__":
    try:  # texto nítido en pantallas con escala en Windows
        import ctypes
        ctypes.windll.shcore.SetProcessDpiAwareness(1)
    except (AttributeError, OSError):
        pass
    App().mainloop()
